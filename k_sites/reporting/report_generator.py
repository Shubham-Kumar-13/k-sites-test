"""
Report Generator for K-Sites

This module generates publication-ready HTML reports from K-Sites pipeline output.
Includes comprehensive CSV and Excel export functionality.
"""

import json
import html
import csv
import os
from typing import Dict, Any, List
from io import StringIO
import logging

from k_sites.reporting.evidence_columns import (
    BASIS_LABELS, EVIDENCE_HEADERS, EXCLUDED_HEADERS, evidence_cells, evidence_summary, excluded_cells,
)

# Set up logging
logger = logging.getLogger(__name__)

CONTEXT_LABELS = {"cells": "cultured cells", "organism": "whole organism"}


def generate_html_report(pipeline_output: dict, output_path: str) -> None:
    """
    Generate a publication-ready HTML report from pipeline output.
    
    Args:
        pipeline_output: Output dictionary from run_k_sites_pipeline
        output_path: Path to save the HTML report
    """
    logger.info(f"Generating HTML report at {output_path}")
    
    try:
        # Extract data from pipeline output
        metadata = pipeline_output.get("metadata", {})
        genes = pipeline_output.get("genes", [])
        insufficient_genes = pipeline_output.get("insufficient_data_genes", [])
        warnings = pipeline_output.get("warnings", [])
        statistics = pipeline_output.get("statistics", {})
        excluded_genes = pipeline_output.get("excluded_genes", [])

        # Count statistics
        total_genes_screened = statistics.get("total_genes_screened", len(genes))

        # Generate the HTML report
        html_content = _generate_report_html(
            metadata,
            genes,
            total_genes_screened,
            output_path,
            insufficient_genes,
            warnings,
            excluded_genes
        )
        
        # Write to file
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html_content)
        
        # Generate additional output formats (CSV, Excel)
        _generate_additional_outputs(pipeline_output, output_path)
        
        logger.info(f"Successfully generated HTML report at {output_path}")
        
    except Exception as e:
        logger.error(f"Failed to generate HTML report: {str(e)}")
        raise


def _generate_additional_outputs(pipeline_output: dict, output_path: str) -> None:
    """
    Generate additional output formats alongside the HTML report.
    
    Args:
        pipeline_output: Output dictionary from run_k_sites_pipeline
        output_path: Path to the HTML report (used to derive other paths)
    """
    try:
        base_path = output_path.replace('.html', '')
        
        # Generate comprehensive CSV
        csv_path = f"{base_path}_comprehensive.csv"
        _generate_csv_report(pipeline_output, csv_path)
        
        # Generate gene summary CSV
        summary_csv_path = f"{base_path}_gene_summary.csv"
        _generate_gene_summary_csv(pipeline_output, summary_csv_path)
        
        # Generate Excel report
        excel_path = f"{base_path}_comprehensive.xlsx"
        _generate_excel_report(pipeline_output, excel_path)
        
        logger.info(f"Generated all additional output formats")
        
    except Exception as e:
        logger.warning(f"Some additional outputs could not be generated: {str(e)}")


def _generate_csv_report(pipeline_output: dict, csv_path: str) -> None:
    """
    Generate a detailed CSV report with all results.
    
    Args:
        pipeline_output: Output dictionary from run_k_sites_pipeline
        csv_path: Path to save the CSV report
    """
    try:
        genes = pipeline_output.get("genes", [])
        metadata = pipeline_output.get("metadata", {})
        
        # Prepare CSV data
        rows = []
        
        # Add metadata as header comments
        rows.append([f"# K-Sites Gene Selection Report"])
        rows.append([f"# GO Term: {metadata.get('go_term', 'N/A')}"])
        rows.append([f"# Organism: {metadata.get('organism', 'N/A')}"])
        rows.append([f"# Timestamp: {metadata.get('timestamp', 'N/A')}"])
        rows.append([f"# Evidence Filter: {metadata.get('evidence_filter', 'N/A')}"])
        rows.append([f"# Pleiotropy cutoff (0-10): {metadata.get('max_pleiotropy', 'N/A')}"])
        data_info = metadata.get("knockout_evidence_data") or {}
        if metadata.get("knockout_context"):
            rows.append([f"# Knockout context: {CONTEXT_LABELS.get(metadata['knockout_context'], metadata['knockout_context'])}"])
        if data_info.get("available"):
            rows.append([f"# Measured knockout data version: {data_info.get('version')}; sources: "
                         + "; ".join(s["name"] for s in data_info.get("sources", []))])
        for warning in pipeline_output.get("warnings", []):
            rows.append([f"# WARNING: {warning}"])
        rows.append(["#"])

        # Add header row
        header = [
            "Gene_Symbol", "Gene_Description", "Entrez_ID", "Pleiotropy_Score",
            "Specificity_Score", "Evidence_Quality", "Literature_Score",
            "Conservation_Score", "Composite_Score", "BP_Term_Count",
            "Experimental_Evidence_Count", "Computational_Evidence_Count",
            "IEA_Evidence_Count", "KEGG_Pathway_Count", "PubMed_Count",
            "Phenotype_Severity", "Phenotype_Risk_Level", "Phenotype_Confidence",
            "Phenotype_Lethality_Stage", "Safety_Level", "Safety_Recommendation",
            "Safety_Justification"
        ] + EVIDENCE_HEADERS
        rows.append(header)
        
        # Add data rows
        for gene in genes:
            gene_symbol = gene.get("symbol", "")
            gene_description = gene.get("description", "")
            entrez_id = gene.get("entrez_id", "")
            pleiotropy_score = gene.get("pleiotropy_score", "")
            specificity_score = gene.get("specificity_score", "")
            evidence_quality = gene.get("evidence_quality", "")
            literature_score = gene.get("literature_score", "")
            conservation_score = gene.get("conservation_score", "")
            composite_score = gene.get("composite_score", "")
            bp_term_count = gene.get("bp_term_count", "")
            exp_count = gene.get("experimental_evidence_count", "")
            comp_count = gene.get("computational_evidence_count", "")
            iea_count = gene.get("iea_evidence_count", "")
            kegg_count = gene.get("kegg_pathway_count", "")
            pubmed_count = gene.get("pubmed_count", "")
            
            # Get phenotype prediction if available
            phenotype_pred = gene.get("phenotype_prediction", {}) or {}
            phenotype_severity = _extract_enum_value(phenotype_pred.get("severity", ""))
            phenotype_risk = _extract_enum_value(phenotype_pred.get("risk_level", ""))
            phenotype_confidence = phenotype_pred.get("confidence_score", "")
            phenotype_stage = phenotype_pred.get("lethality_stage", "")
            
            # Get safety recommendation
            safety_rec = gene.get("safety_recommendation", {}) or {}
            safety_level = safety_rec.get("safety_level", "")
            primary_rec = safety_rec.get("primary_recommendation", "")
            safety_justification = _excel_text(safety_rec.get("justification", ""))

            rows.append([
                gene_symbol, gene_description, entrez_id, pleiotropy_score,
                specificity_score, evidence_quality, literature_score,
                conservation_score, composite_score, bp_term_count,
                exp_count, comp_count, iea_count, kegg_count, pubmed_count,
                phenotype_severity, phenotype_risk, phenotype_confidence,
                phenotype_stage, safety_level, primary_rec, safety_justification
            ] + evidence_cells(gene))

        # Genes that were assessed but are not in the table above
        excluded = pipeline_output.get("excluded_genes", [])
        if excluded:
            rows.append([])
            rows.append(["# Genes not shown above (above the pleiotropy cutoff, or ranked below the genes shown)"])
            rows.append(EXCLUDED_HEADERS)
            rows.extend(excluded_cells(entry) for entry in excluded)

        # Write CSV file
        with open(csv_path, 'w', newline='', encoding='utf-8') as csvfile:
            writer = csv.writer(csvfile)
            writer.writerows(rows)
        
        logger.info(f"Generated CSV report at {csv_path}")
        
    except Exception as e:
        logger.error(f"Failed to generate CSV report: {str(e)}")
        raise


def _generate_gene_summary_csv(pipeline_output: dict, csv_path: str) -> None:
    """
    Generate a gene-focused summary CSV report.
    
    Args:
        pipeline_output: Output dictionary from run_k_sites_pipeline
        csv_path: Path to save the CSV report
    """
    try:
        genes = pipeline_output.get("genes", [])
        
        with open(csv_path, 'w', newline='', encoding='utf-8') as csvfile:
            fieldnames = [
                'gene_symbol', 'gene_description', 'entrez_id', 
                'pleiotropy_score', 'specificity_score', 'composite_score',
                'evidence_quality', 'literature_score', 'conservation_score',
                'bp_term_count', 'experimental_evidence_count', 
                'computational_evidence_count', 'iea_evidence_count',
                'kegg_pathway_count', 'pubmed_count',
                'phenotype_severity', 'phenotype_risk_level',
                'safety_level', 'safety_recommendation', 'safety_justification'
            ] + [name.lower() for name in EVIDENCE_HEADERS]
            writer = csv.DictWriter(csvfile, fieldnames=fieldnames)
            
            writer.writeheader()
            for gene in genes:
                phenotype_pred = gene.get("phenotype_prediction", {}) or {}
                safety_rec = gene.get("safety_recommendation", {}) or {}
                
                writer.writerow({
                    'gene_symbol': gene.get('symbol', ''),
                    'gene_description': gene.get('description', ''),
                    'entrez_id': gene.get('entrez_id', ''),
                    'pleiotropy_score': gene.get('pleiotropy_score', ''),
                    'specificity_score': gene.get('specificity_score', ''),
                    'composite_score': gene.get('composite_score', ''),
                    'evidence_quality': gene.get('evidence_quality', ''),
                    'literature_score': gene.get('literature_score', ''),
                    'conservation_score': gene.get('conservation_score', ''),
                    'bp_term_count': gene.get('bp_term_count', ''),
                    'experimental_evidence_count': gene.get('experimental_evidence_count', ''),
                    'computational_evidence_count': gene.get('computational_evidence_count', ''),
                    'iea_evidence_count': gene.get('iea_evidence_count', ''),
                    'kegg_pathway_count': gene.get('kegg_pathway_count', ''),
                    'pubmed_count': gene.get('pubmed_count', ''),
                    'phenotype_severity': _extract_enum_value(phenotype_pred.get('severity', '')),
                    'phenotype_risk_level': _extract_enum_value(phenotype_pred.get('risk_level', '')),
                    'safety_level': safety_rec.get('safety_level', ''),
                    'safety_recommendation': safety_rec.get('primary_recommendation', ''),
                    'safety_justification': _excel_text(safety_rec.get('justification', '')),
                    **{name.lower(): value for name, value in zip(EVIDENCE_HEADERS, evidence_cells(gene))}
                })
        
        logger.info(f"Generated gene summary CSV at {csv_path}")
        
    except Exception as e:
        logger.error(f"Failed to generate gene summary CSV: {str(e)}")


def _generate_excel_report(pipeline_output: dict, excel_path: str) -> None:
    """
    Generate a comprehensive Excel report with multiple sheets.
    
    Args:
        pipeline_output: Output dictionary from run_k_sites_pipeline
        excel_path: Path to save the Excel report
    """
    try:
        # Try to import openpyxl for Excel generation
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Fill, PatternFill, Alignment, Border, Side
            from openpyxl.utils.dataframe import dataframe_to_rows
            from openpyxl.utils import get_column_letter
        except ImportError:
            logger.warning("openpyxl not available, generating CSV fallback instead of Excel")
            # Generate CSV as fallback
            csv_fallback = excel_path.replace('.xlsx', '_excel_fallback.csv')
            _generate_csv_report(pipeline_output, csv_fallback)
            return
        
        genes = pipeline_output.get("genes", [])
        metadata = pipeline_output.get("metadata", {})
        statistics = pipeline_output.get("statistics", {})
        
        wb = Workbook()
        
        # ===== Sheet 1: Summary =====
        ws_summary = wb.active
        ws_summary.title = "Summary"
        
        # Style definitions
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        
        # Add metadata
        ws_summary['A1'] = "K-Sites Gene Selection Report"
        ws_summary['A1'].font = Font(bold=True, size=16)
        ws_summary.merge_cells('A1:D1')
        
        ws_summary['A3'] = "GO Term:"
        ws_summary['B3'] = metadata.get('go_term', 'N/A')
        ws_summary['A4'] = "Organism:"
        ws_summary['B4'] = metadata.get('organism', 'N/A')
        ws_summary['A5'] = "Timestamp:"
        ws_summary['B5'] = metadata.get('timestamp', 'N/A')
        ws_summary['A6'] = "Evidence Filter:"
        ws_summary['B6'] = metadata.get('evidence_filter', 'N/A')
        ws_summary['A7'] = "Pleiotropy cutoff (0-10):"
        ws_summary['B7'] = metadata.get('max_pleiotropy', 'N/A')
        
        # Add statistics
        ws_summary['A9'] = "Statistics"
        ws_summary['A9'].font = Font(bold=True, size=12)
        ws_summary['A10'] = "Total Genes Screened:"
        ws_summary['B10'] = statistics.get('total_genes_screened', 0)
        ws_summary['A11'] = "Genes Passed Filter:"
        ws_summary['B11'] = statistics.get('genes_passed_filter', 0)
        ws_summary['A12'] = "Average Pleiotropy:"
        ws_summary['B12'] = f"{statistics.get('avg_pleiotropy', 0):.2f}"
        summary_row = 13
        if metadata.get("knockout_context"):
            ws_summary.cell(row=summary_row, column=1, value="Knockout context:")
            ws_summary.cell(row=summary_row, column=2, value=CONTEXT_LABELS.get(metadata["knockout_context"], metadata["knockout_context"]))
            summary_row += 1
        for warning in pipeline_output.get("warnings", []):
            ws_summary.cell(row=summary_row, column=1, value="WARNING:")
            ws_summary.cell(row=summary_row, column=2, value=warning)
            summary_row += 1
        
        # ===== Sheet 2: Gene Rankings =====
        ws_genes = wb.create_sheet("Gene Rankings")
        
        # Headers
        gene_headers = [
            "Rank", "Gene Symbol", "Pleiotropy Score", "Specificity Score",
            "Composite Score", "Evidence Quality", "Literature Score",
            "Conservation Score", "BP Term Count", "Safety Level", "Evidence Basis"
        ]
        
        for col, header in enumerate(gene_headers, 1):
            cell = ws_genes.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        
        # Data
        for row, gene in enumerate(genes, 2):
            ws_genes.cell(row=row, column=1, value=row-1)  # Rank
            ws_genes.cell(row=row, column=2, value=gene.get('symbol', ''))
            ws_genes.cell(row=row, column=3, value=gene.get('pleiotropy_score', ''))
            ws_genes.cell(row=row, column=4, value=gene.get('specificity_score', ''))
            ws_genes.cell(row=row, column=5, value=gene.get('composite_score', ''))
            ws_genes.cell(row=row, column=6, value=gene.get('evidence_quality', ''))
            ws_genes.cell(row=row, column=7, value=gene.get('literature_score', ''))
            ws_genes.cell(row=row, column=8, value=gene.get('conservation_score', ''))
            ws_genes.cell(row=row, column=9, value=gene.get('bp_term_count', ''))
            
            safety_rec = gene.get('safety_recommendation', {}) or {}
            ws_genes.cell(row=row, column=10, value=safety_rec.get('safety_level', ''))
            ws_genes.cell(row=row, column=11, value=safety_rec.get('evidence_basis', ''))
        
        # ===== Sheet 3: Safety Recommendations =====
        ws_safety = wb.create_sheet("Safety Recommendations")
        
        safety_headers = [
            "Gene Symbol", "Safety Level", "Primary Recommendation",
            "Justification", "Concerns", "Mitigation Strategies"
        ]
        
        for col, header in enumerate(safety_headers, 1):
            cell = ws_safety.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        
        for row, gene in enumerate(genes, 2):
            safety_rec = gene.get('safety_recommendation', {}) or {}
            ws_safety.cell(row=row, column=1, value=gene.get('symbol', ''))
            ws_safety.cell(row=row, column=2, value=safety_rec.get('safety_level', ''))
            ws_safety.cell(row=row, column=3, value=safety_rec.get('primary_recommendation', ''))
            ws_safety.cell(row=row, column=4, value=_excel_text(safety_rec.get('justification', '')))
            ws_safety.cell(row=row, column=5, value='; '.join(safety_rec.get('concerns', [])))
            ws_safety.cell(row=row, column=6, value='; '.join(safety_rec.get('mitigation_strategies', [])))
        
        # ===== Sheet 4: Measured knockout evidence =====
        ws_evidence = wb.create_sheet("Knockout Evidence")
        for col, header in enumerate(["Gene Symbol"] + EVIDENCE_HEADERS, 1):
            cell = ws_evidence.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        for row, gene in enumerate(genes, 2):
            for col, value in enumerate([gene.get('symbol', '')] + evidence_cells(gene), 1):
                ws_evidence.cell(row=row, column=col, value=_excel_text(value))

        # ===== Sheet 5: Genes assessed but not shown =====
        sheets = [ws_summary, ws_genes, ws_safety, ws_evidence]
        excluded = pipeline_output.get("excluded_genes", [])
        if excluded:
            ws_excluded = wb.create_sheet("Genes Not Shown")
            for col, header in enumerate(EXCLUDED_HEADERS, 1):
                cell = ws_excluded.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
            for row, entry in enumerate(excluded, 2):
                for col, value in enumerate(excluded_cells(entry), 1):
                    ws_excluded.cell(row=row, column=col, value=_excel_text(value))
            sheets.append(ws_excluded)

        # Auto-adjust column widths
        for ws in sheets:
            for column in ws.columns:
                max_length = 0
                # column[0] can be a merged cell (title row), which has no column_letter
                column_letter = get_column_letter(column[0].column)
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save workbook
        wb.save(excel_path)
        logger.info(f"Generated Excel report at {excel_path}")
        
    except Exception as e:
        logger.error(f"Failed to generate Excel report: {str(e)}")


SAFETY_CLASSES = {
    "SAFE": "safety-safe",                  # green
    "CAUTION": "safety-caution",            # yellow
    "MODERATE_CONCERN": "safety-concern",   # orange
    "HIGH_CONCERN": "safety-warning",       # red
    "AVOID": "safety-critical",             # dark red
}


def _safety_class(level) -> str:
    """CSS class for a safety level (grey for anything unrecognised)."""
    return SAFETY_CLASSES.get(str(level or "").upper(), "safety-unknown")


def _safety_label(level) -> str:
    """Readable label for a safety level, e.g. MODERATE_CONCERN -> 'Moderate concern'."""
    text = str(level or "").strip()
    if not text:
        return "Unknown"
    return text.replace("_", " ").capitalize()


def _excel_text(value) -> str:
    """Excel cells hold scalars: join lists into one string."""
    if isinstance(value, (list, tuple, set)):
        return "; ".join(str(v) for v in value)
    return value if value is not None else ""


def _extract_enum_value(value) -> str:
    """Extract string value from enum or return as-is if string."""
    if hasattr(value, 'value'):
        return str(value.value)
    return str(value) if value else ""


def _evidence_html(entry: dict) -> str:
    """The measured knockout evidence of a gene (or of a gene left out of the ranking) as report HTML."""
    evidence = entry.get("knockout_evidence")
    basis = (entry.get("safety_recommendation") or {}).get("evidence_basis")
    if basis is None and evidence:
        # genes left out of the ranking carry no verdict, so the basis is read from the evidence itself
        basis = "none"
        if evidence["has_measured"]:
            basis = "measured" if evidence["organism"]["basis"] == "direct" or evidence["cells"]["call"] != "UNKNOWN" else "ortholog"
        elif evidence["organism"]["call"] != "UNKNOWN" and evidence["organism"]["distant"]:
            basis = "distant"
        elif evidence.get("estimate"):
            basis = "estimate"
    basis = basis or "none"
    parts = [f'<span class="basis basis-{html.escape(basis)}">{html.escape(BASIS_LABELS.get(basis, basis))}</span>']
    if evidence:
        organism, cells = evidence["organism"], evidence["cells"]
        for part in (organism, cells):
            if part["call"] != "UNKNOWN":
                parts.append(f'<div class="evidence-line">{html.escape(part["text"])}</div>')
        if evidence.get("estimate") and organism["call"] == "UNKNOWN":
            parts.append(f'<div class="evidence-line">{html.escape(evidence["estimate"]["text"])}</div>')
        links = [f'<a href="https://pubmed.ncbi.nlm.nih.gov/{int(p)}/">PMID {int(p)}</a>' for p in organism["references"]]
        if organism["mgi_id"]:
            links.append(f'<a href="https://www.informatics.jax.org/marker/{html.escape(organism["mgi_id"])}">MGI</a>')
        if links:
            parts.append(f'<div class="evidence-refs">{" | ".join(links)}</div>')
    return "".join(parts)


def _generate_report_html(
    metadata: Dict[str, Any], 
    genes: list, 
    total_genes_screened: int,
    output_path: str,
    insufficient_genes: list = None,
    warnings: list = None,
    excluded_genes: list = None
) -> str:
    """
    Generate the complete HTML report.
    """
    excluded_genes = excluded_genes or []
    # Escape metadata values for safety
    escaped_go_term = html.escape(str(metadata.get("go_term", "Unknown")))
    escaped_organism = html.escape(str(metadata.get("organism", "Unknown")))
    resolved_organism = metadata.get("resolved_organism", {})
    escaped_scientific_name = html.escape(str(resolved_organism.get("scientific_name", "Unknown") if resolved_organism else "Unknown"))
    timestamp = metadata.get("timestamp", "")
    execution_duration = metadata.get("execution_duration", 0)
    evidence_filter = metadata.get("evidence_filter", "experimental")
    max_pleiotropy = metadata.get("max_pleiotropy", 10)
    
    insufficient_genes = insufficient_genes or []
    knockout_context = metadata.get("knockout_context")
    context_html = (f" | Knockout context: <strong>{html.escape(CONTEXT_LABELS.get(knockout_context, str(knockout_context)))}</strong>"
                    if knockout_context else "")
    warnings_html = "".join(
        f'<div class="report-warning"><strong>Warning:</strong> {html.escape(str(w))}</div>'
        for w in (warnings or [])
    )

    # Safety tallies for the executive summary
    def _level(g):
        return (g.get("safety_recommendation", {}) or {}).get("safety_level", "")
    safe_count = sum(1 for g in genes if _level(g) == "SAFE")
    caution_count = sum(1 for g in genes if _level(g) in ("CAUTION", "MODERATE_CONCERN"))
    concern_count = sum(1 for g in genes if _level(g) in ("HIGH_CONCERN", "AVOID"))
    
    # Build the HTML content
    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>K-Sites Gene Selection Report - {escaped_go_term}</title>
    <style>
        :root {{
            --primary-color: #2c3e50;
            --secondary-color: #3498db;
            --success-color: #27ae60;
            --warning-color: #f39c12;
            --danger-color: #e74c3c;
            --light-bg: #f8f9fa;
            --border-color: #dee2e6;
            --text-color: #212529;
            --text-light: #6c757d;
        }}
        
        * {{ margin: 0; padding: 0; box-sizing: border-box; }}
        
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            line-height: 1.6;
            color: var(--text-color);
            background-color: white;
            padding: 20px;
        }}
        
        .container {{ max-width: 1200px; margin: 0 auto; }}
        
        header {{
            text-align: center;
            margin-bottom: 30px;
            padding-bottom: 20px;
            border-bottom: 2px solid var(--border-color);
        }}
        
        h1 {{ color: var(--primary-color); margin-bottom: 10px; font-size: 2.2em; }}
        
        .subtitle {{ color: var(--text-light); font-size: 1.1em; margin-bottom: 15px; }}
        
        .summary-card {{
            background: linear-gradient(135deg, #f5f7fa 0%, #e4edf9 100%);
            border-radius: 10px;
            padding: 20px;
            margin: 20px 0;
            box-shadow: 0 4px 6px rgba(0,0,0,0.1);
            border-left: 4px solid var(--secondary-color);
        }}
        
        .summary-grid {{
            display: grid;
            grid-template-columns: repeat(auto-fit, minmax(200px, 1fr));
            gap: 20px;
            margin-top: 15px;
        }}
        
        .summary-item {{
            text-align: center;
            padding: 15px;
            background: white;
            border-radius: 8px;
            box-shadow: 0 2px 4px rgba(0,0,0,0.05);
        }}
        
        .summary-value {{ font-size: 1.8em; font-weight: bold; color: var(--primary-color); }}
        .summary-label {{ font-size: 0.9em; color: var(--text-light); }}
        
        .section {{ margin: 30px 0; }}
        
        .section-title {{
            color: var(--primary-color);
            border-bottom: 2px solid var(--secondary-color);
            padding-bottom: 10px;
            margin-bottom: 20px;
            font-size: 1.5em;
        }}
        
        table {{
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 20px;
            background: white;
            box-shadow: 0 2px 8px rgba(0,0,0,0.1);
            border-radius: 8px;
            overflow: hidden;
        }}
        
        th, td {{ padding: 12px 15px; text-align: left; border-bottom: 1px solid var(--border-color); }}
        th {{ background-color: var(--primary-color); color: white; font-weight: 600; }}
        tr:nth-child(even) {{ background-color: #f8f9fa; }}
        tr:hover {{ background-color: #e9f7fe; }}
        
        
        .score-high {{ color: var(--success-color); font-weight: bold; }}
        .score-medium {{ color: var(--warning-color); font-weight: bold; }}
        .score-low {{ color: var(--danger-color); font-weight: bold; }}
        
        
        .phenotype-severity, .risk-level, .safety-level {{
            display: inline-block;
            padding: 4px 8px;
            border-radius: 12px;
            font-size: 0.8em;
            font-weight: bold;
            text-transform: uppercase;
        }}
        
        .safety-safe {{ background-color: #d4edda; color: #155724; }}
        .safety-caution {{ background-color: #fff3cd; color: #856404; }}
        .safety-warning {{ background-color: #f8d7da; color: #721c24; }}
        .safety-concern {{ background-color: #ffe5d0; color: #8a4b08; }}
        .safety-critical {{ background-color: #dc3545; color: white; }}
        .safety-unknown {{ background-color: #e2e3e5; color: #41464b; }}
        
        .report-warning {{
            background-color: #fff3cd;
            color: #664d03;
            border-left: 4px solid #f39c12;
            border-radius: 6px;
            padding: 12px 16px;
            margin: 20px 0;
        }}
        
        .risk-critical {{ background-color: #dc3545; color: white; }}
        .risk-high {{ background-color: #fd7e14; color: white; }}
        .risk-medium {{ background-color: #ffc107; color: #212529; }}
        .risk-low {{ background-color: #28a745; color: white; }}
        .risk-unknown {{ background-color: #6c757d; color: white; }}
        
        .visual-indicator {{
            display: inline-block;
            width: 12px;
            height: 12px;
            border-radius: 50%;
            margin-right: 5px;
        }}
        
        .high-specificity {{ background-color: #28a745; }}
        .medium-specificity {{ background-color: #ffc107; }}
        .low-specificity {{ background-color: #dc3545; }}
        
        .specificity-bar {{
            height: 10px;
            background-color: #e9ecef;
            border-radius: 5px;
            overflow: hidden;
            margin: 5px 0;
            width: 100px;
        }}
        
        .specificity-fill {{ height: 100%; background-color: #28a745; }}
        
        .download-section {{
            text-align: center;
            margin: 40px 0;
            padding: 30px;
            background: linear-gradient(135deg, #e8f4f8 0%, #d4edf9 100%);
            border-radius: 10px;
        }}
        
        .download-btn {{
            display: inline-block;
            background-color: var(--secondary-color);
            color: white;
            padding: 12px 24px;
            text-decoration: none;
            border-radius: 8px;
            font-weight: 600;
            margin: 10px;
            transition: all 0.3s;
        }}
        .download-btn:hover {{
            background-color: #2980b9;
            transform: translateY(-2px);
        }}
        
        .justification-box {{
            background-color: #f8f9fa;
            padding: 10px;
            border-radius: 4px;
            font-size: 0.9em;
            margin-top: 5px;
            border-left: 3px solid var(--secondary-color);
        }}

        .basis {{ display: inline-block; padding: 2px 8px; border-radius: 10px; font-size: 0.75em; font-weight: bold; margin-bottom: 4px; }}
        .basis-measured {{ background: #d4edda; color: #155724; }}
        .basis-ortholog {{ background: #d1ecf1; color: #0c5460; }}
        .basis-predicted {{ background: #fff3cd; color: #856404; }}
        .basis-estimate {{ background: #fff3cd; color: #856404; }}
        .basis-distant {{ background: #e2e3e5; color: #41464b; }}
        .basis-none {{ background: #e2e3e5; color: #41464b; }}
        .evidence-line {{ font-size: 0.88em; margin-top: 3px; }}
        .evidence-refs {{ font-size: 0.8em; color: var(--text-light); margin-top: 3px; }}
        .sources-list {{ font-size: 0.9em; color: var(--text-light); }}

        @media (max-width: 768px) {{
            body {{ padding: 10px; }}
            .summary-grid {{ grid-template-columns: 1fr; }}
            th, td {{ padding: 8px 10px; font-size: 0.9em; }}
        }}
    </style>
</head>
<body>
    <div class="container">
        <header>
            <h1>🧬 K-Sites Gene Selection Report</h1>
            <div class="subtitle">
                GO Term: <strong>{escaped_go_term}</strong> | 
                Organism: <strong>{escaped_scientific_name}</strong> | 
                Generated: {timestamp}
            </div>
            <div class="subtitle">
                Evidence Filter: <strong>{evidence_filter}</strong> |
                Pleiotropy cutoff (0-10): <strong>{max_pleiotropy}</strong>{context_html}
            </div>
        </header>

        {warnings_html}

        <div class="summary-card">
            <h2>📊 Executive Summary</h2>
            <div class="summary-grid">
                <div class="summary-item">
                    <div class="summary-value">{total_genes_screened}</div>
                    <div class="summary-label">Genes Screened</div>
                </div>
                <div class="summary-item">
                    <div class="summary-value">{safe_count}</div>
                    <div class="summary-label">Safe to Delete</div>
                </div>
                <div class="summary-item">
                    <div class="summary-value">{caution_count}</div>
                    <div class="summary-label">Proceed with Caution</div>
                </div>
                <div class="summary-item">
                    <div class="summary-value">{concern_count}</div>
                    <div class="summary-label">High Concern / Avoid</div>
                </div>
                <div class="summary-item">
                    <div class="summary-value">{len(insufficient_genes)}</div>
                    <div class="summary-label">Not Scored (Insufficient Data)</div>
                </div>
            </div>
        </div>
        
        <div class="section">
            <h2 class="section-title">🧬 Gene Rankings with Specificity Indicators</h2>
            <p style="margin-bottom: 15px; color: var(--text-light);">
                Genes ranked by composite score combining specificity (40%), evidence quality (25%), 
                literature support (20%), and conservation (15%). Specificity score is on 0-1 scale 
                (higher = more specific, less pleiotropic).
            </p>
            <table>
                <thead>
                    <tr>
                        <th>Rank</th>
                        <th>Gene Symbol</th>
                        <th>Pleiotropy (0-10)</th>
                        <th>Specificity (0-1)</th>
                        <th>Visual</th>
                        <th>Composite Score</th>
                        <th>Safety Level</th>
                        <th>Measured Knockout Evidence</th>
                        <th>KEGG Pathways</th>
                        <th>Recommendation</th>
                    </tr>
                </thead>
                <tbody>
"""
    
    # Add gene rows to the table
    for rank, gene in enumerate(genes, 1):
        symbol = html.escape(str(gene.get("symbol", "Unknown")))
        pleiotropy_score = gene.get("pleiotropy_score", 0)
        specificity_score = gene.get("specificity_score", 0)
        composite_score = gene.get("composite_score", 0)
        
        # Ensure specificity is 0-1 scale
        if specificity_score > 1.0:
            specificity_score = specificity_score / 10.0
        
        specificity_percentage = specificity_score * 100
        
        # Get safety recommendation
        safety_rec = gene.get("safety_recommendation", {}) or {}
        safety_level = safety_rec.get("safety_level", "UNKNOWN")
        primary_rec = safety_rec.get("primary_recommendation", "")
        safety_justification = safety_rec.get("justification", "")
        
        # Determine visual indicator class
        if specificity_score >= 0.7:
            indicator_class = "high-specificity"
        elif specificity_score >= 0.4:
            indicator_class = "medium-specificity"
        else:
            indicator_class = "low-specificity"
        
        # Determine safety class
        safety_class = _safety_class(safety_level)
        
        kegg_count = gene.get("kegg_pathway_count", "N/A")
        if isinstance(safety_justification, (list, tuple)):
            safety_justification = "; ".join(str(line) for line in safety_justification)
        evidence_html = _evidence_html(gene)

        html_content += f"""
                    <tr>
                        <td>{rank}</td>
                        <td><strong>{symbol}</strong></td>
                        <td>{pleiotropy_score:.2f}</td>
                        <td>{specificity_score:.3f}</td>
                        <td>
                            <div class="visual-indicator {indicator_class}"></div>
                            <div class="specificity-bar">
                                <div class="specificity-fill" style="width: {specificity_percentage}%"></div>
                            </div>
                        </td>
                        <td>{composite_score:.3f}</td>
                        <td><span class="safety-level {safety_class}">{html.escape(_safety_label(safety_level))}</span></td>
                        <td>{evidence_html}</td>
                        <td>{kegg_count}</td>
                        <td>
                            {primary_rec}
                            {f'<div class="justification-box">{html.escape(str(safety_justification))}</div>' if safety_justification else ''}
                        </td>
                    </tr>
"""
    
    html_content += """
                </tbody>
            </table>
        </div>
        
"""
    
    if insufficient_genes:
        html_content += """
        <div class="section">
            <h2 class="section-title">⚠️ Genes Not Scored (Insufficient Data)</h2>
            <p style="margin-bottom: 15px; color: var(--text-light);">
                These genes could not be assessed, so they are not ranked and no safety level is given.
                "Lookup failed" means a database could not be reached; the gene may be fine on retry.
            </p>
            <table>
                <thead>
                    <tr><th>Gene Symbol</th><th>Reason</th><th>Lookup Failed</th></tr>
                </thead>
                <tbody>
"""
        for gene in insufficient_genes:
            reasons = "; ".join(str(r) for r in gene.get("reasons", [])) or "No data"
            html_content += f"""
                    <tr>
                        <td><strong>{html.escape(str(gene.get("symbol", "")))}</strong></td>
                        <td>{html.escape(reasons)}</td>
                        <td>{"Yes" if gene.get("lookup_failed") else "No"}</td>
                    </tr>
"""
        html_content += """
                </tbody>
            </table>
        </div>
"""

    if excluded_genes:
        html_content += """
        <div class="section">
            <h2 class="section-title">Genes Assessed but Not Shown</h2>
            <p style="margin-bottom: 15px; color: var(--text-light);">
                These genes were assessed too, but are not in the ranking above: either their pleiotropy score is above the
                cutoff, or they rank below the genes shown. Their measured knockout result is listed so that nothing is hidden.
            </p>
            <table>
                <thead>
                    <tr><th>Gene Symbol</th><th>Why Not Shown</th><th>Measured Knockout Evidence</th></tr>
                </thead>
                <tbody>
"""
        for entry in excluded_genes:
            html_content += f"""
                    <tr>
                        <td><strong>{html.escape(str(entry.get("symbol", "")))}</strong></td>
                        <td>{html.escape(str(entry.get("detail", "")))}</td>
                        <td>{_evidence_html(entry)}</td>
                    </tr>
"""
        html_content += """
                </tbody>
            </table>
        </div>
"""

    data_info = metadata.get("knockout_evidence_data") or {}
    if data_info.get("sources"):
        items = "".join(
            f'<li><a href="{html.escape(s["url"])}">{html.escape(s["name"])}</a>: {html.escape(s.get("what", ""))}</li>'
            for s in data_info["sources"])
        html_content += f"""
        <div class="section">
            <h2 class="section-title">Data Sources</h2>
            <p class="sources-list">Measured knockout data version {html.escape(str(data_info.get("version")))}.</p>
            <ul class="sources-list">{items}</ul>
        </div>
"""

    # Get base filename for download links
    base_name = os.path.basename(output_path).replace('.html', '')
    
    html_content += f"""
        <div class="download-section">
            <h2>📥 Download Results</h2>
            <p style="margin-bottom: 20px; color: var(--text-light);">
                Download results in various formats for further analysis.
            </p>
            <a href="./{base_name}_comprehensive.csv" class="download-btn">📄 Comprehensive CSV</a>
            <a href="./{base_name}_gene_summary.csv" class="download-btn">📊 Gene Summary CSV</a>
            <a href="./{base_name}_comprehensive.xlsx" class="download-btn">📗 Excel Report</a>
        </div>
        
        <footer style="text-align: center; margin-top: 40px; padding-top: 20px; border-top: 1px solid var(--border-color); color: var(--text-light); font-size: 0.9em;">
            <p>K-Sites Gene Selection Platform | Non-Pleiotropic Gene Identification System</p>
            <p>Execution time: {execution_duration:.2f}s | Evidence: {evidence_filter} | Pleiotropy cutoff (0-10): {max_pleiotropy}</p>
        </footer>
    </div>
    
</body>
</html>"""
    
    return html_content
